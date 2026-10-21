#!/usr/bin/env python3
"""Build 'CryptAlyst — a domain model' as an A4 landscape PDF.

Cover is a full-bleed plate with no page number; the body carries the number in the
footer. Every section opens with what it shows, then the figure, then the argument.
"""
import os

from PIL import Image
from docx import Document
from docx.enum.section import WD_ORIENT, WD_SECTION
from docx.enum.table import WD_TABLE_ALIGNMENT
from docx.enum.text import WD_ALIGN_PARAGRAPH, WD_BREAK
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Cm, Emu, Inches, Pt, RGBColor

import content

HERE = os.path.dirname(os.path.abspath(__file__))
FIG = os.path.join(HERE, "fig")
ACCENT = RGBColor(0x1B, 0x4A, 0x6F)
BRIGHT = RGBColor(0x2E, 0x6D, 0xB4)
INK = RGBColor(0x18, 0x20, 0x2A)
MUTED = RGBColor(0x5C, 0x66, 0x72)
FAINT = RGBColor(0x8A, 0x94, 0xA0)
BAND = "EEF3F8"

PAGE_W, PAGE_H = Cm(29.7), Cm(21.0)          # landscape A4
MARGIN_X, MARGIN_Y = Inches(0.85), Inches(0.65)
CONTENT_W = 29.7 / 2.54 - 1.7                 # ≈ 10.0 in

doc = Document()
cover_sec = doc.sections[0]
cover_sec.orientation = WD_ORIENT.LANDSCAPE
cover_sec.page_width, cover_sec.page_height = PAGE_W, PAGE_H
for side in ("left_margin", "right_margin", "top_margin", "bottom_margin"):
    setattr(cover_sec, side, Emu(0))

normal = doc.styles["Normal"]
normal.font.name = "Calibri"
normal.font.size = Pt(10)
normal.font.color.rgb = INK
normal.paragraph_format.space_after = Pt(6)
normal.paragraph_format.line_spacing = 1.12


def shade(paragraph, colour, left_bar=None):
    pr = paragraph._p.get_or_add_pPr()
    shd = OxmlElement("w:shd")
    shd.set(qn("w:val"), "clear"); shd.set(qn("w:fill"), colour)
    pr.append(shd)
    if left_bar:
        borders = OxmlElement("w:pBdr")
        left = OxmlElement("w:left")
        left.set(qn("w:val"), "single"); left.set(qn("w:sz"), "24")
        left.set(qn("w:space"), "8"); left.set(qn("w:color"), left_bar)
        borders.append(left); pr.append(borders)


def h1(text, number=None, new_page=False):
    para = doc.add_paragraph()
    if new_page:
        para.paragraph_format.page_break_before = True
    para.paragraph_format.space_before = Pt(2)
    para.paragraph_format.space_after = Pt(6)
    if number:
        r = para.add_run("%s   " % number)
        r.font.size = Pt(11); r.bold = True; r.font.color.rgb = FAINT
    r = para.add_run(text)
    r.font.size = Pt(19); r.bold = True; r.font.color.rgb = ACCENT


def shows(text):
    para = doc.add_paragraph()
    para.paragraph_format.space_before = Pt(0)
    para.paragraph_format.space_after = Pt(8)
    shade(para, BAND, "2E6DB4")
    r = para.add_run("WHAT THIS SHOWS   ")
    r.font.size = Pt(7.5); r.bold = True; r.font.color.rgb = BRIGHT
    r = para.add_run(text)
    r.font.size = Pt(9.5); r.font.color.rgb = INK


def para(text, size=10, italic=False, colour=None, space=6):
    p = doc.add_paragraph()
    p.paragraph_format.space_after = Pt(space)
    r = p.add_run(text)
    r.font.size = Pt(size); r.italic = italic
    if colour: r.font.color.rgb = colour
    return p


def bullets(items, size=9.6):
    for it in items:
        p = doc.add_paragraph(style="List Bullet")
        p.paragraph_format.space_after = Pt(2)
        p.paragraph_format.left_indent = Inches(0.22)
        p.add_run(it).font.size = Pt(size)


def table(head, rows, size=8.8):
    t = doc.add_table(rows=1, cols=len(head))
    t.style = "Table Grid"
    t.alignment = WD_TABLE_ALIGNMENT.CENTER
    for i, name in enumerate(head):
        cell = t.rows[0].cells[i]
        cell.text = ""
        r = cell.paragraphs[0].add_run(name.upper())
        r.bold = True; r.font.size = Pt(7.4); r.font.color.rgb = MUTED
        shd = OxmlElement("w:shd"); shd.set(qn("w:fill"), "F1F4F7")
        cell._tc.get_or_add_tcPr().append(shd)
    for row in rows:
        cells = t.add_row().cells
        for i, val in enumerate(row):
            cells[i].text = ""
            run = cells[i].paragraphs[0].add_run(val)
            run.font.size = Pt(size)
    # no trailing paragraph: it is what pushes a blank page when a table ends a page


def figure(name, caption, max_h=3.95):
    path = os.path.join(FIG, name)
    with Image.open(path) as im:
        w, h = im.size
    width = min(CONTENT_W, 10.0)
    height = width * h / w
    if height > max_h:
        height = max_h
        width = height * w / h
    doc.add_picture(path, width=Inches(width))
    doc.paragraphs[-1].alignment = WD_ALIGN_PARAGRAPH.CENTER
    doc.paragraphs[-1].paragraph_format.space_after = Pt(2)
    cap = doc.add_paragraph()
    cap.alignment = WD_ALIGN_PARAGRAPH.CENTER
    cap.paragraph_format.space_after = Pt(6)
    r = cap.add_run(caption)
    r.italic = True; r.font.size = Pt(8.2); r.font.color.rgb = FAINT


def two_columns(left_text, right_items):
    t = doc.add_table(rows=1, cols=2)
    t.autofit = False
    widths = (Inches(6.1), Inches(3.8))
    for i, cell in enumerate(t.rows[0].cells):
        cell.width = widths[i]
        cell.text = ""
    left = t.rows[0].cells[0].paragraphs[0]
    left.add_run(left_text).font.size = Pt(9.8)
    right = t.rows[0].cells[1]
    first = True
    for item in right_items:
        p = right.paragraphs[0] if first else right.add_paragraph()
        first = False
        p.paragraph_format.space_after = Pt(2)
        run = p.add_run("· ")
        run.font.color.rgb = BRIGHT; run.font.size = Pt(9.4)
        p.add_run(item).font.size = Pt(9.4)


# ------------------------------------------------------------------ cover
cover = doc.add_paragraph()
cover.paragraph_format.space_after = Pt(0)
run = cover.add_run()
run.add_picture(os.path.join(FIG, "00-cover.png"), width=PAGE_W, height=PAGE_H)
cover.alignment = WD_ALIGN_PARAGRAPH.CENTER

# ------------------------------------------------------------------ body section
body_sec = doc.add_section(WD_SECTION.NEW_PAGE)
body_sec.orientation = WD_ORIENT.LANDSCAPE
body_sec.page_width, body_sec.page_height = PAGE_W, PAGE_H
body_sec.left_margin = body_sec.right_margin = MARGIN_X
body_sec.top_margin = body_sec.bottom_margin = MARGIN_Y
body_sec.footer.is_linked_to_previous = False

# the cover must not carry a page number
cover_sec.footer.is_linked_to_previous = False
cover_sec.footer.paragraphs[0].text = ""

footer = body_sec.footer.paragraphs[0]
footer.alignment = WD_ALIGN_PARAGRAPH.CENTER
r = footer.add_run("CryptAlyst — a domain model   ·   ")
r.font.size = Pt(7.5); r.font.color.rgb = FAINT
fld = OxmlElement("w:fldSimple"); fld.set(qn("w:instr"), "PAGE")
footer._p.append(fld)

# ------------------------------------------------------------------ contents
h1("Contents", "")
para("Every section begins with what it shows; every figure is listed here and captioned "
     "where it appears. The document runs from the whole platform down to the filesystem.",
     size=9.8, colour=MUTED, space=10)

sections = content.SECTIONS
half = (len(sections) + 1) // 2
t = doc.add_table(rows=1, cols=2)
t.autofit = False
for i, cell in enumerate(t.rows[0].cells):
    cell.width = Inches(5.0)
    cell.text = ""

for col, chunk in ((0, sections[:half]), (1, sections[half:])):
    cell = t.rows[0].cells[col]
    first = True
    for s in chunk:
        p = cell.paragraphs[0] if first else cell.add_paragraph()
        first = False
        p.paragraph_format.space_after = Pt(5)
        num = s.get("number") or "—"
        r = p.add_run("%-5s " % num); r.font.size = Pt(8.5); r.font.color.rgb = FAINT
        r = p.add_run(s["title"]); r.font.size = Pt(9.6); r.bold = True; r.font.color.rgb = INK
        p2 = cell.add_paragraph()
        p2.paragraph_format.left_indent = Inches(0.28)
        p2.paragraph_format.space_after = Pt(5)
        p2.add_run(s.get("shows", "")).font.size = Pt(8.4)
        for run_ in p2.runs:
            run_.font.color.rgb = MUTED

p = doc.add_paragraph()
p.paragraph_format.page_break_before = True      # the index gets its own page
p.paragraph_format.space_before = Pt(0)
p.paragraph_format.space_after = Pt(4)
r = p.add_run("FIGURES")
r.font.size = Pt(8); r.bold = True; r.font.color.rgb = FAINT

rows = [[n, title, what] for n, _f, title, what in content.FIGURES if n != "00"]
table(["Fig", "Figure", "What it shows"], rows, size=7.8)

# ------------------------------------------------------------------ sections
for n, s in enumerate(sections, 1):
    h1(s["title"], s.get("number") or "", new_page=True)
    shows(s.get("shows", ""))
    if s.get("figure"):
        name, caption = s["figure"]
        fig_no = next((f[0] for f in content.FIGURES if f[1] == name), "")
        height = 3.15 if s.get("table") else 3.7
        figure(name, "Figure %s — %s" % (fig_no, caption), max_h=height)
    body = s.get("body", [])
    if body and s.get("bullets"):
        two_columns(body[0], s["bullets"])
        for text in body[1:]:
            para(text)
    else:
        for text in body:
            para(text)
        if s.get("bullets"):
            bullets(s["bullets"])
    if s.get("table"):
        head, rows = s["table"]
        table(head, rows)
    for text in s.get("body2", []):
        para(text)

# ------------------------------------------------------------------ colophon
h1("Colophon", "", new_page=True)
shows("Where this document comes from, and how to regenerate it.")
para(content.VERSION + "   ·   " + content.REVISION, size=9.6, colour=INK)
para(content.COPYRIGHT, size=9.2, colour=MUTED)
para("The text lives once in content.py; the Markdown and this PDF are both generated from it, so "
     "they cannot disagree. The figures are drawn as SVG by diagrams.py and cover.py, rendered "
     "to PNG by render_figs.py, and placed here by build_pdf.py.", size=9.4)
table(["Artifact", "Command"], [
    ["figures", "python3 diagrams.py && python3 cover.py && python3 render_figs.py"],
    ["markdown", "python3 build_md.py"],
    ["pdf", "python3 build_pdf.py && libreoffice --convert-to pdf"],
])

out = os.path.join(HERE, "CryptAlyst-Domain-Model.docx")
doc.save(out)
print("saved", out)
